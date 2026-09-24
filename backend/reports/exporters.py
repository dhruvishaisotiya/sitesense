import csv
import io
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors


class ReportExporterService:

    @classmethod
    def export_csv(cls, report_type, data):
        response = HttpResponse(content_type='text/csv')
        filename = f"SiteSense_{report_type}_Report.csv"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'

        writer = csv.writer(response)

        if report_type == 'projects' and isinstance(data, list):
            writer.writerow(['Project Code', 'Project Name', 'Building Type', 'Manager', 'Status', 'Current Stage', 'Progress %', 'Budget Used ($)', 'Total Budget ($)', 'Delay Prob %', 'Risk', 'Recommendation'])
            for item in data:
                writer.writerow([
                    item.get('project_code', ''),
                    item.get('project_name', ''),
                    item.get('building_type', ''),
                    item.get('manager', ''),
                    item.get('status', ''),
                    item.get('current_stage', ''),
                    item.get('current_progress', 0.0),
                    item.get('budget_used', 0.0),
                    item.get('total_budget', 0.0),
                    item.get('delay_probability', 0.0),
                    item.get('project_risk', ''),
                    item.get('recommendation', '')
                ])
        elif report_type in ['ai', 'ai-prediction'] and isinstance(data, list):
            writer.writerow(['Project Code/ID', 'Delay Probability %', 'Completion Days Remaining', 'Risk Level', 'AI Suggestion', 'Prediction Time'])
            for item in data:
                writer.writerow([
                    item.get('project_code', f"Project {item.get('project_id', '')}"),
                    item.get('delay_probability', 0.0),
                    item.get('completion_days_remaining', 0),
                    item.get('project_risk', item.get('risk', '')),
                    item.get('ai_suggestion', item.get('recommendation', '')),
                    item.get('prediction_timestamp', '')
                ])
        elif report_type == 'dailylogs' and isinstance(data, dict):
            writer.writerow(['Day Number', 'Project Code', 'Stage', 'Progress %', 'Attendance %', 'Rainfall (mm)'])
            for item in data.get('progress_trend', []):
                writer.writerow([
                    item.get('day_number', ''),
                    item.get('project_code', ''),
                    item.get('construction_stage', ''),
                    item.get('progress_percentage', 0.0),
                    item.get('attendance_percentage', 0.0),
                    item.get('rainfall_mm', 0.0)
                ])
        elif report_type == 'workers' and isinstance(data, dict):
            writer.writerow(['Metric / Project', 'Value / Active Workers'])
            writer.writerow(['Total Workers', data.get('total_workers', 0)])
            writer.writerow(['Permanent', data.get('permanent_count', data.get('permanent', 0))])
            writer.writerow(['Contract', data.get('contract_count', data.get('contract', 0))])
            writer.writerow(['Daily Wage', data.get('daily_wage_count', data.get('daily_wage', 0))])
            writer.writerow(['Active Workers', data.get('active_count', data.get('active', 0))])
            writer.writerow(['Inactive Workers', data.get('inactive_count', data.get('inactive', 0))])
            for wp in data.get('workers_per_project', []):
                writer.writerow([f"{wp.get('project_code')} - {wp.get('project_name')}", wp.get('active_workers', 0)])
        elif report_type == 'tasks' and isinstance(data, dict):
            writer.writerow(['Task Metric / Priority', 'Count'])
            writer.writerow(['Pending Tasks', data.get('pending', 0)])
            writer.writerow(['In Progress Tasks', data.get('in_progress', 0)])
            writer.writerow(['Completed Tasks', data.get('completed', 0)])
            writer.writerow(['Approved Tasks', data.get('approved', 0)])
            writer.writerow(['Rejected Tasks', data.get('rejected', 0)])
            writer.writerow(['Overdue Tasks', data.get('overdue', 0)])
            for bp in data.get('by_priority', []):
                writer.writerow([f"{bp.get('priority')} Priority", bp.get('count', 0)])
        elif report_type == 'budget' and isinstance(data, dict):
            writer.writerow(['Metric / Expense Category', 'Amount ($)'])
            writer.writerow(['Total Budget', data.get('total_budget', data.get('project_budget', 0.0))])
            writer.writerow(['Budget Used', data.get('budget_used', 0.0)])
            writer.writerow(['Remaining Budget', data.get('remaining_budget', 0.0)])
            writer.writerow(['Material Expenses', data.get('material_expenses', 0.0)])
            writer.writerow(['Budget Utilization %', data.get('utilization_pct', data.get('budget_utilization_pct', 0.0))])
            for cat in data.get('expense_breakdown', []):
                writer.writerow([cat.get('category', ''), cat.get('amount', 0.0)])
        else:
            writer.writerow(['Report Title', f'SiteSense {report_type.capitalize()} Report'])
            if isinstance(data, dict):
                for k, v in data.items():
                    if isinstance(v, (int, float, str)):
                        writer.writerow([k, v])

        return response

    @classmethod
    def export_excel(cls, report_type, data):
        wb = Workbook()
        ws = wb.active
        ws.title = f"{report_type.capitalize()} Report"

        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='003135', end_color='003135', fill_type='solid')

        if report_type == 'projects' and isinstance(data, list):
            headers = ['Project Code', 'Project Name', 'Building Type', 'Manager', 'Status', 'Current Stage', 'Progress %', 'Budget Used ($)', 'Total Budget ($)', 'Delay Prob %', 'Risk', 'Recommendation']
            ws.append(headers)
            for item in data:
                ws.append([
                    item.get('project_code', ''),
                    item.get('project_name', ''),
                    item.get('building_type', ''),
                    item.get('manager', ''),
                    item.get('status', ''),
                    item.get('current_stage', ''),
                    item.get('current_progress', 0.0),
                    item.get('budget_used', 0.0),
                    item.get('total_budget', 0.0),
                    item.get('delay_probability', 0.0),
                    item.get('project_risk', ''),
                    item.get('recommendation', '')
                ])
        elif report_type in ['ai', 'ai-prediction'] and isinstance(data, list):
            headers = ['Project Code/ID', 'Delay Prob %', 'Completion Days Rem.', 'Risk Level', 'AI Suggestion', 'Timestamp']
            ws.append(headers)
            for item in data:
                ws.append([
                    item.get('project_code', f"Project {item.get('project_id', '')}"),
                    item.get('delay_probability', 0.0),
                    item.get('completion_days_remaining', 0),
                    item.get('project_risk', item.get('risk', '')),
                    item.get('ai_suggestion', item.get('recommendation', '')),
                    item.get('prediction_timestamp', '')
                ])
        elif report_type == 'dailylogs' and isinstance(data, dict):
            headers = ['Day Number', 'Project Code', 'Stage', 'Progress %', 'Attendance %', 'Rainfall (mm)']
            ws.append(headers)
            for item in data.get('progress_trend', []):
                ws.append([
                    item.get('day_number', ''),
                    item.get('project_code', ''),
                    item.get('construction_stage', ''),
                    item.get('progress_percentage', 0.0),
                    item.get('attendance_percentage', 0.0),
                    item.get('rainfall_mm', 0.0)
                ])
        else:
            headers = ['Metric / Key', 'Value']
            ws.append(headers)
            if isinstance(data, dict):
                for k, v in data.items():
                    if isinstance(v, (int, float, str)):
                        ws.append([k.replace('_', ' ').title(), v])

        # Style headers
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(output.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        filename = f"SiteSense_{report_type}_Report.xlsx"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    @classmethod
    def export_pdf(cls, report_type, data):
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'ReportTitle',
            parent=styles['Heading1'],
            fontSize=16,
            textColor=colors.HexColor('#003135'),
            spaceAfter=8
        )
        subtitle_style = ParagraphStyle(
            'ReportSubTitle',
            parent=styles['Normal'],
            fontSize=9,
            textColor=colors.HexColor('#0FA4AF'),
            spaceAfter=15
        )

        elements.append(Paragraph(f"SiteSense — Executive {report_type.capitalize()} Report", title_style))
        elements.append(Paragraph("Generated directly from live system database records.", subtitle_style))

        if report_type == 'projects' and isinstance(data, list):
            table_data = [['Code', 'Project Name', 'Stage', 'Progress', 'Budget Used', 'Risk', 'AI Suggestion']]
            for item in data:
                b_used = int(float(item.get('budget_used', 0) or 0))
                table_data.append([
                    str(item.get('project_code', '') or ''),
                    Paragraph(str(item.get('project_name', '') or ''), styles['Normal']),
                    str(item.get('current_stage', '') or ''),
                    f"{item.get('current_progress', 0)}%",
                    f"${b_used:,}",
                    str(item.get('project_risk', '') or ''),
                    Paragraph(str(item.get('recommendation', '') or ''), styles['Normal'])
                ])
            t = Table(table_data, colWidths=[65, 120, 80, 50, 75, 45, 115])
        elif report_type in ['ai', 'ai-prediction'] and isinstance(data, list):
            table_data = [['Project', 'Delay Prob', 'Days Rem.', 'Risk', 'AI Suggestion']]
            for item in data:
                table_data.append([
                    str(item.get('project_code', f"Project {item.get('project_id', '')}")),
                    f"{item.get('delay_probability', 0)}%",
                    f"{item.get('completion_days_remaining', 0)} d",
                    str(item.get('project_risk', item.get('risk', 'Low'))),
                    Paragraph(str(item.get('ai_suggestion', item.get('recommendation', ''))), styles['Normal'])
                ])
            t = Table(table_data, colWidths=[80, 70, 70, 60, 220])
        elif report_type == 'dailylogs' and isinstance(data, dict):
            table_data = [['Day', 'Project Code', 'Stage', 'Progress %', 'Attendance %', 'Rainfall']]
            for item in data.get('progress_trend', []):
                table_data.append([
                    f"Day {item.get('day_number', '')}",
                    str(item.get('project_code', '')),
                    str(item.get('construction_stage', '')),
                    f"{item.get('progress_percentage', 0)}%",
                    f"{item.get('attendance_percentage', 0)}%",
                    f"{item.get('rainfall_mm', 0)} mm"
                ])
            t = Table(table_data, colWidths=[60, 100, 140, 70, 75, 55])
        else:
            table_data = [['Metric Key', 'Value']]
            if isinstance(data, dict):
                for k, v in data.items():
                    if isinstance(v, (int, float, str)):
                        table_data.append([k.replace('_', ' ').title(), str(v)])

            t = Table(table_data, colWidths=[250, 250])

        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#003135')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 9),
            ('BOTTOMPADDING', (0,0), (-1,0), 8),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ]))
        elements.append(t)

        doc.build(elements)
        buffer.seek(0)

        response = HttpResponse(buffer.read(), content_type='application/pdf')
        filename = f"SiteSense_{report_type}_Report.pdf"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
